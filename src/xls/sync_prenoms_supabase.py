#!/usr/bin/env python3
"""
Synchronisation complète du catalogue « Fête des prénoms » depuis le fichier
Excel fusionné vers :
  1. les fichiers de données locaux (src/data/prenoms-{filles,garcons,mixtes}.{json,ts})
  2. les images générées (public/catalogue/prenoms/{filles,garcons,mixtes})
  3. la table Supabase public.prenoms (genres : fille, garcon, mixte)

Source : 37x37_Prenoms_fusion.11.xlsm
  - « Prènom Fille »  : Prénom | Description | date texte | date datetime
  - « Prènom Garçon » : Prénom | date texte | date datetime
  - mixte = prénoms présents dans les DEUX listes (vrais unisexes)

Note : la colonne « Sexe = Mixte » de l'onglet « Prénoms avec fête » ne désigne
pas un prénom unisexe mais un prénom fêté pour les deux genres — elle n'est
pas utilisée pour la classification affichée sur le site.

Usage : python3 sync_prenoms_supabase.py
"""

import json
import os
import sys
import unicodedata
from pathlib import Path

import openpyxl
import urllib.request
import urllib.error
from PIL import Image, ImageDraw, ImageFont

EXCEL_PATH = "/Users/stephanevigneron/Desktop/Marché de Noel/37x37_Prenoms_fusion.11.xlsm"
PROJECT = Path("/Users/stephanevigneron/Desktop/site stephane/mon-aimant-pr-cieux")

import os

def _load_env(path):
    if not os.path.exists(path):
        return
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

_load_env(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".env"))

SUPABASE_URL = os.environ.get("SUPABASE_URL", "https://fcjxacmdzhzjdiqjsbid.supabase.co")
SERVICE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
HEADERS = {
    "apikey": SERVICE_KEY,
    "Authorization": f"Bearer {SERVICE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}

MONTHS = {
    "janvier": 1, "février": 2, "fevrier": 2, "mars": 3, "avril": 4,
    "mai": 5, "juin": 6, "juillet": 7, "août": 8, "aout": 8,
    "septembre": 9, "octobre": 10, "novembre": 11, "décembre": 12, "decembre": 12,
}

GENRES = {
    "fille": {"dir": "filles", "prefix": "prenom-fille", "color": "#F7B4C0", "ts": "PrenomFille", "const": "PRENOMS_FILLES"},
    "garcon": {"dir": "garcons", "prefix": "prenom-garcon", "color": "#9FC5E8", "ts": "PrenomGarcon", "const": "PRENOMS_GARCONS"},
    "mixte": {"dir": "mixtes", "prefix": "prenom-mixte", "color": "#C9A7EB", "ts": "PrenomMixte", "const": "PRENOMS_MIXTES"},
}

IMG_SIZE = 600


def norm_key(s: str) -> str:
    return unicodedata.normalize("NFD", s).casefold()


def parse_label_date(label):
    """« le 15 Avril » → « 2024-04-15 »."""
    if not label:
        return None
    parts = label.strip().lower().split()
    if len(parts) != 3 or parts[0] != "le":
        return None
    try:
        day = int(parts[1])
    except ValueError:
        return None
    month = MONTHS.get(parts[2])
    if not month:
        return None
    return f"2024-{month:02d}-{day:02d}"


def read_excel():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

    # Onglet fille : Prénom | Description | date texte | date datetime
    filles = {}  # prenom -> (label, date)
    for r in wb["Prènom Fille"].iter_rows(min_row=2, values_only=True):
        if r[0] and str(r[0]).strip():
            label = (str(r[2]).strip() or None) if r[2] else None
            filles.setdefault(str(r[0]).strip(), (label, norm_date(r[3]) or parse_label_date(label)))

    # Onglet garçon : Prénom | date texte | date datetime
    garcons = {}
    for r in wb["Prènom Garçon"].iter_rows(values_only=True):
        if r[0] and str(r[0]).strip() and str(r[0]).strip().lower() != "prénom":
            label = (str(r[1]).strip() or None) if r[1] else None
            garcons.setdefault(str(r[0]).strip(), (label, norm_date(r[2]) or parse_label_date(label)))

    return filles, garcons


def norm_date(raw):
    """Datetime Excel → '2024-MM-DD' (année normalisée)."""
    if raw and hasattr(raw, "month") and hasattr(raw, "day"):
        return f"2024-{raw.month:02d}-{raw.day:02d}"
    return None


def build_datasets(filles, garcons):
    """Retourne {genre: [(prenom, fete_label, fete_date, sexe)]} trié alphabétiquement."""
    communs = set(filles) & set(garcons)

    data = {
        "fille": [
            (p, label, fdate, "Mixte" if p in communs else "Fille")
            for p, (label, fdate) in filles.items()
        ],
        "garcon": [
            (p, label, fdate, "Mixte" if p in communs else "Garçon")
            for p, (label, fdate) in garcons.items()
        ],
        "mixte": [],
    }
    for p in communs:
        (ff, fd), (fg, _) = filles[p], garcons[p]
        label = ff if ff == fg else f"{ff} / {fg}"
        data["mixte"].append((p, label, fd, "Mixte"))

    for g in data:
        data[g].sort(key=lambda t: norm_key(t[0]))
    return data


def write_data_files(data):
    for genre, rows in data.items():
        meta = GENRES[genre]
        payload = [{"prenom": p, "fete_label": label, "sexe": sexe} for p, label, _, sexe in rows]

        json_path = PROJECT / f"src/data/prenoms-{meta['dir']}.json"
        json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

        body = ",\n".join(
            '  {\n    "prenom": %s,\n    "fete_label": %s,\n    "sexe": %s\n  }'
            % (json.dumps(p, ensure_ascii=False), json.dumps(label, ensure_ascii=False), json.dumps(sexe, ensure_ascii=False))
            for p, label, _, sexe in rows
        )
        ts = (
            f"export interface {meta['ts']} {{\n"
            "  prenom: string;\n"
            "  fete_label: string | null;\n"
            "  sexe: string | null;\n"
            "}\n\n"
            f"export const {meta['const']}: {meta['ts']}[] = [\n{body},\n];\n"
        )
        (PROJECT / f"src/data/prenoms-{meta['dir']}.ts").write_text(ts, encoding="utf-8")
        print(f"  {genre}: {len(rows)} entrées → {json_path.name} / .ts")


def find_font():
    for p in (
        "/System/Library/Fonts/Supplemental/ChalkboardSE.ttc",
        "/System/Library/Fonts/Supplemental/Chalkboard.ttc",
        "/System/Library/Fonts/MarkerFelt.ttc",
        "/System/Library/Fonts/Helvetica.ttc",
    ):
        if os.path.exists(p):
            return p
    return None


def fit_text(draw, text, max_width, font_path, max_size=120, min_size=40):
    for size in range(max_size, min_size - 1, -4):
        font = ImageFont.truetype(font_path, size)
        bbox = draw.textbbox((0, 0), text, font=font)
        if bbox[2] - bbox[0] <= max_width:
            return font, bbox[2] - bbox[0], bbox[3] - bbox[1]
    font = ImageFont.truetype(font_path, min_size)
    bbox = draw.textbbox((0, 0), text, font=font)
    return font, bbox[2] - bbox[0], bbox[3] - bbox[1]


def make_image(prenom, fete_label, color, font_path):
    img = Image.new("RGB", (IMG_SIZE, IMG_SIZE), color)
    draw = ImageDraw.Draw(img)
    prenom_font, pw, ph = fit_text(draw, prenom, IMG_SIZE - 80, font_path)

    parts = []
    total_h = ph
    if fete_label:
        sf_font = ImageFont.truetype(font_path, 36)
        sf_box = draw.textbbox((0, 0), "se fête", font=sf_font)
        d_font = ImageFont.truetype(font_path, 50)
        d_box = draw.textbbox((0, 0), fete_label, font=d_font)
        total_h += 110 + (sf_box[3] - sf_box[1]) + 75 + (d_box[3] - d_box[1])
        parts = [(sf_font, "se fête", sf_box), (d_font, fete_label, d_box)]

    y = (IMG_SIZE - total_h) // 2
    draw.text(((IMG_SIZE - pw) // 2, y), prenom, font=prenom_font, fill="#1a1a1a")
    if parts:
        sf_font, _, sf_box = parts[0]
        d_font, d_text, d_box = parts[1]
        y += ph + 110
        draw.text(((IMG_SIZE - (sf_box[2] - sf_box[0])) // 2, y), "se fête", font=sf_font, fill="#1a1a1a")
        y += (sf_box[3] - sf_box[1]) + 75
        draw.text(((IMG_SIZE - (d_box[2] - d_box[0])) // 2, y), d_text, font=d_font, fill="#1a1a1a")
    return img


def generate_images(data):
    font_path = find_font()
    if not font_path:
        print("  ⚠️  aucune police trouvée, images ignorées")
        return
    for genre, rows in data.items():
        meta = GENRES[genre]
        out_dir = PROJECT / "public/catalogue/prenoms" / meta["dir"]
        out_dir.mkdir(parents=True, exist_ok=True)
        # purge des anciens fichiers du préfixe
        for old in out_dir.glob(f"{meta['prefix']}-*.png"):
            old.unlink()
        for i, (prenom, label, *_rest) in enumerate(rows):
            img = make_image(prenom, label, meta["color"], font_path)
            img.save(out_dir / f"{meta['prefix']}-{1000 + i}.png")
        print(f"  {genre}: {len(rows)} images → {out_dir.relative_to(PROJECT)}")


def rest(method, path, payload=None):
    req = urllib.request.Request(
        f"{SUPABASE_URL}/rest/v1/{path}",
        method=method,
        headers=HEADERS,
        data=json.dumps(payload).encode() if payload is not None else None,
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            body = r.read().decode()
            return r.status, json.loads(body) if body else None
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


def sync_supabase(data):
    print("\nSuppression des anciennes lignes…")
    status, _ = rest("DELETE", "prenoms?id=gt.0")
    print(f"  DELETE → {status}")

    total = 0
    for genre, rows in data.items():
        meta = GENRES[genre]
        payload = [
            {
                "prenom": p,
                "genre": genre,
                "fete_label": label,
                "fete_date": fdate,
                "photo_url": f"/catalogue/prenoms/{meta['dir']}/{meta['prefix']}-{1000 + i}.png",
                "active": True,
            }
            for i, (p, label, fdate, _sexe) in enumerate(rows)
        ]
        inserted = 0
        for j in range(0, len(payload), 500):
            status, resp = rest("POST", "prenoms", payload[j : j + 500])
            if status not in (200, 201):
                print(f"  ✗ {genre} chunk {j}: HTTP {status} {resp}")
                sys.exit(1)
            inserted += len(resp) if isinstance(resp, list) else len(payload[j : j + 500])
        print(f"  {genre}: {inserted} lignes insérées")
        total += inserted
    print(f"  Total : {total} lignes")


def main():
    data_only = "--data-only" in sys.argv

    print("Lecture du fichier Excel…")
    filles, garcons = read_excel()
    print(f"  {len(filles)} filles, {len(garcons)} garçons, {len(set(filles) & set(garcons))} mixtes")

    data = build_datasets(filles, garcons)
    print("\nÉcriture des fichiers data…")
    write_data_files(data)

    if data_only:
        print("\n✅ Terminé (data uniquement, images et Supabase inchangés).")
        return

    print("\nGénération des images…")
    generate_images(data)

    print("\nSynchronisation Supabase…")
    sync_supabase(data)

    print("\n✅ Terminé.")


if __name__ == "__main__":
    main()
