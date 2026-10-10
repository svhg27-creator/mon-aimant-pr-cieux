#!/usr/bin/env python3
"""
Extrait les noms de pays du fichier 37x37 Drapeau coeur .xlsx
et met à jour la colonne pays dans public.catalogue_drapeaux.
Usage : python3 update_drapeaux_noms.py
"""

import openpyxl
import psycopg2

XLSX_PATH = "/Users/stephanevigneron/Desktop/site stephane/mon-aimant-pr-cieux/src/xls/37x37 Drapeau coeur .xlsx"
DB_URL = "postgresql://postgres:1HVEqHSfAlQlonUc@db.fcjxacmdzhzjdiqjsbid.supabase.co:5432/postgres"


def extract_noms():
    """Extrait les noms de pays dans l'ordre visuel du fichier Excel."""
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    noms = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        # Les noms sont sous les images, aux lignes 4, 7, 10, 13, 16
        # et aux colonnes B (2), D (4), F (6)
        for row in [4, 7, 10, 13, 16]:
            for col in [2, 4, 6]:
                val = ws.cell(row=row, column=col).value
                if val and str(val).strip():
                    text = str(val).strip()
                    # Ignorer les cellules d'instruction/titre
                    if "🖼️" in text or "🌍" in text or text.lower().startswith("page"):
                        continue
                    noms.append(text)
    return noms


def update_database(noms):
    """Met à jour les noms de pays en base."""
    print(f"Noms extraits : {len(noms)}")

    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()

    # Compter le nombre d'images en base
    cur.execute("SELECT COUNT(*) FROM public.catalogue_drapeaux WHERE active = true")
    count = cur.fetchone()[0]
    print(f"Images en base : {count}")

    # On associe les noms aux images par ordre d'index
    limit = min(len(noms), count)
    updated = 0
    for i, nom in enumerate(noms[:limit], 1):
        cur.execute(
            "UPDATE public.catalogue_drapeaux SET pays = %s WHERE image_index = %s",
            (nom, i)
        )
        if cur.rowcount > 0:
            updated += 1

    conn.commit()
    cur.close()
    conn.close()
    print(f"\n✅ {updated} noms de pays mis à jour")


def main():
    noms = extract_noms()
    update_database(noms)


if __name__ == "__main__":
    main()
