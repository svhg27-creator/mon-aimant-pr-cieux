#!/usr/bin/env python3
"""
Crée un fichier Excel de vérification avec toutes les images de drapeaux
et une colonne pour corriger les noms de pays.
Usage : python3 create_verification_excel.py
"""

import psycopg2
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenPyXLImage
from pathlib import Path

DB_URL = "postgresql://postgres:1HVEqHSfAlQlonUc@db.fcjxacmdzhzjdiqjsbid.supabase.co:5432/postgres"
OUTPUT_XLSX = "/Users/stephanevigneron/Desktop/site stephane/mon-aimant-pr-cieux/src/xls/drapeaux_verification.xlsx"
IMAGE_DIR = "/Users/stephanevigneron/Desktop/site stephane/mon-aimant-pr-cieux/public/catalogue/drapeaux"


def main():
    print("Connexion à la base...")
    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()
    cur.execute("SELECT image_index, pays, image_name FROM public.catalogue_drapeaux ORDER BY image_index")
    rows = cur.fetchall()
    cur.close()
    conn.close()

    print(f"{len(rows)} drapeaux trouvés")

    wb = Workbook()
    ws = wb.active
    ws.title = "Vérification"

    # En-têtes
    ws.cell(row=1, column=1, value="Index")
    ws.cell(row=1, column=2, value="Image")
    ws.cell(row=1, column=3, value="Nom actuel")
    ws.cell(row=1, column=4, value="Nom corrigé")

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 30

    for idx, (image_index, pays, image_name) in enumerate(rows, start=2):
        ws.cell(row=idx, column=1, value=image_index)
        ws.cell(row=idx, column=3, value=pays)

        image_path = Path(IMAGE_DIR) / image_name
        if image_path.exists():
            img = OpenPyXLImage(str(image_path))
            img.width = 100
            img.height = 100
            ws.add_image(img, f"B{idx}")

        ws.row_dimensions[idx].height = 80

    wb.save(OUTPUT_XLSX)
    print(f"\n✅ Fichier de vérification créé : {OUTPUT_XLSX}")
    print("Instructions :")
    print("  1. Ouvre ce fichier Excel")
    print("  2. Regarde chaque image")
    print("  3. Écris le bon nom de pays dans la colonne D (Nom corrigé)")
    print("  4. Sauvegarde le fichier")
    print("  5. Dis-moi 'c'est corrigé' pour que je réimporte")


if __name__ == "__main__":
    main()
