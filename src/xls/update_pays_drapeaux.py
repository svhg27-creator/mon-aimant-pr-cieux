#!/usr/bin/env python3
"""
Extrait les noms de pays du fichier Excel selon la position exacte des images
et met à jour uniquement la colonne 'pays' dans public.catalogue_drapeaux.
"""

import zipfile
import xml.etree.ElementTree as ET
import openpyxl
import psycopg2
from pathlib import Path
import re

XLSX_PATH = "/Users/stephanevigneron/Desktop/site stephane/mon-aimant-pr-cieux/src/xls/37x37 Drapeau coeur .xlsx"
DB_URL = "postgresql://postgres:1HVEqHSfAlQlonUc@db.fcjxacmdzhzjdiqjsbid.supabase.co:5432/postgres"

NS = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}


def parse_drawing(content):
    root = ET.fromstring(content)
    images = []
    for anchor in root.findall(".//xdr:twoCellAnchor", NS):
        pic = anchor.find(".//xdr:pic", NS)
        if pic is None:
            continue
        blip = pic.find(".//a:blip", {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"})
        if blip is None:
            continue
        rId = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        if not rId:
            continue
        from_cell = anchor.find("xdr:from", NS)
        col = int(from_cell.find("xdr:col", NS).text)
        row = int(from_cell.find("xdr:row", NS).text)
        images.append({"rId": rId, "col": col, "row": row})
    return images


def extract_noms_by_position():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    with zipfile.ZipFile(XLSX_PATH) as z:
        drawing_files = sorted(
            [f for f in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", f)],
            key=lambda x: int(re.search(r"\d+", x).group())
        )

        noms = []
        for drawing_path in drawing_files:
            drawing_name = Path(drawing_path).name
            drawing_num = int(re.search(r"\d+", drawing_name).group())
            ws = wb[wb.sheetnames[drawing_num - 1]]
            images = parse_drawing(z.read(drawing_path).decode("utf-8"))
            images.sort(key=lambda x: (x["row"], x["col"]))
            for img in images:
                name_col = (img["col"] // 2) * 2 + 1
                best_nom = None
                for delta in [1, 2, 3]:
                    val = ws.cell(row=img["row"] + delta, column=name_col + 1).value
                    if val and str(val).strip():
                        text = str(val).strip()
                        if "🖼️" not in text and "🌍" not in text and not text.lower().startswith("page"):
                            best_nom = text
                            break
                noms.append(best_nom or f"Drapeau {len(noms) + 1}")
    return noms


def update_database(noms):
    print(f"Noms extraits : {len(noms)}")
    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()

    updated = 0
    for i, nom in enumerate(noms, 1):
        cur.execute(
            "UPDATE public.catalogue_drapeaux SET pays = %s WHERE image_index = %s",
            (nom, i)
        )
        if cur.rowcount > 0:
            updated += 1

    conn.commit()
    cur.close()
    conn.close()
    print(f"✅ {updated} noms de pays mis à jour")


if __name__ == "__main__":
    noms = extract_noms_by_position()
    update_database(noms)
