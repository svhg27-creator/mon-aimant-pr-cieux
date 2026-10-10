#!/usr/bin/env python3
"""
Importe les corrections de noms depuis drapeaux_verification.xlsx
et met à jour la base de données.
Usage : python3 import_corrections.py
"""

import psycopg2
from openpyxl import load_workbook

DB_URL = "postgresql://postgres:1HVEqHSfAlQlonUc@db.fcjxacmdzhzjdiqjsbid.supabase.co:5432/postgres"
VERIFICATION_XLSX = "/Users/stephanevigneron/Desktop/site stephane/mon-aimant-pr-cieux/src/xls/drapeaux_verification.xlsx"


def main():
    print("Lecture du fichier de correction...")
    wb = load_workbook(VERIFICATION_XLSX, data_only=True)
    ws = wb["Vérification"]

    corrections = []
    for row in range(2, ws.max_row + 1):
        index = ws.cell(row=row, column=1).value
        nom_corrige = ws.cell(row=row, column=4).value

        if index is not None and nom_corrige and str(nom_corrige).strip():
            corrections.append((int(index), str(nom_corrige).strip()))

    print(f"{len(corrections)} corrections trouvées")

    if corrections:
        print("\nMise à jour de la base de données...")
        conn = psycopg2.connect(DB_URL)
        cur = conn.cursor()

        updated = 0
        for image_index, pays in corrections:
            cur.execute(
                "UPDATE public.catalogue_drapeaux SET pays = %s WHERE image_index = %s",
                (pays, image_index)
            )
            if cur.rowcount > 0:
                updated += 1
                print(f"  ✅ {image_index:03d} → {pays}")

        conn.commit()
        cur.close()
        conn.close()
        print(f"\n✅ {updated} noms mis à jour")
    else:
        print("Aucune correction trouvée dans la colonne D.")


if __name__ == "__main__":
    main()
